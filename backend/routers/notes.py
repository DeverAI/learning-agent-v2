"""笔记/考点 API 路由"""

import io
import asyncio
import base64
import binascii
import os
import re
import json
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete, func, or_, TEXT
from models.database import get_db, async_session
from models.models import Note, Question, gen_id
from services.tag_unification_service import canonicalize_tags
from services.note_service import resolve_note_references
from logger import get_logger, log_error
from pydantic import BaseModel, Field, field_validator, ConfigDict
from datetime import datetime, timezone

def _utcnow():
    return datetime.now(timezone.utc).replace(tzinfo=None)
from config import STORAGE_DIR

logger = get_logger()
router = APIRouter(prefix="/api/notes", tags=["notes"])

MAX_BASE64_PER_IMAGE = 15 * 1024 * 1024  # ~10MB raw image → ~13.3MB base64
MAX_IMAGES = 10
MAX_CAPTURE_TOTAL_BYTES = 60 * 1024 * 1024  # 单次请求解码后图片总量上限，防 150MB 级 body 全量进内存

# 笔记入库/合并全程串行化：读快照→AI 合并→写回之间有多个 await，
# 并发上传两批笔记会基于过期快照互相覆盖（丢内容/丢图）。
_persist_note_lock = asyncio.Lock()


class NoteUploadRequest(BaseModel):
    images: list[str] = Field(max_length=MAX_IMAGES)  # base64编码的图片列表
    subject_hint: str = Field(default="", max_length=64)  # 学科提示（可选）

    @field_validator("images")
    @classmethod
    def validate_images(cls, v):
        if len(v) > MAX_IMAGES:
            raise ValueError(f"最多上传 {MAX_IMAGES} 张图片")
        total_decoded = 0
        for i, img in enumerate(v):
            if len(img) > MAX_BASE64_PER_IMAGE:
                raise ValueError(f"第 {i+1} 张图片过大（{len(img)//1024}KB），单张限制 {MAX_BASE64_PER_IMAGE//1024//1024}MB")
            if "," in img:
                header, payload = img.split(",", 1)
                if not header.lower().startswith(("data:image/jpeg", "data:image/png", "data:image/webp")):
                    raise ValueError(f"第 {i+1} 张仅支持 JPG、PNG 或 WEBP")
            else:
                payload = img
            try:
                raw = base64.b64decode(payload, validate=True)
            except (binascii.Error, ValueError):
                raise ValueError(f"第 {i+1} 张图片编码无效")
            valid_signature = (raw.startswith(b"\xff\xd8\xff") or raw.startswith(b"\x89PNG\r\n\x1a\n")
                               or (len(raw) >= 12 and raw[:4] == b"RIFF" and raw[8:12] == b"WEBP"))
            if not valid_signature:
                raise ValueError(f"第 {i+1} 张图片内容无效")
            # 总量上限：逐张累计，防止 10×15MB 合法单张拼出超大请求撑爆内存
            total_decoded += len(raw)
            if total_decoded > MAX_CAPTURE_TOTAL_BYTES:
                raise ValueError(f"图片总大小超过 {MAX_CAPTURE_TOTAL_BYTES//1024//1024}MB 限制，请分批上传")
        return v


class NoteUpdateRequest(BaseModel):
    title: str | None = Field(default=None, max_length=200)
    content: str | None = Field(default=None, max_length=500_000)
    knowledge_tags: list[str] | None = Field(default=None, max_length=100)
    diagram_spec: dict | None = None
    subject: str | None = Field(default=None, max_length=64)
    grade: str | None = Field(default=None, max_length=64)

    @field_validator("diagram_spec")
    @classmethod
    def validate_diagram_spec(cls, value):
        if value is not None and len(json.dumps(value, ensure_ascii=False, default=str)) > 200_000:
            raise ValueError("笔记图形规格不能超过 200KB")
        return value

    model_config = ConfigDict(extra="forbid")


class NoteReferencesRequest(BaseModel):
    references: list[dict] = Field(default_factory=list, max_length=100)


class NoteCreateRequest(BaseModel):
    """手动创建笔记请求（用于「另存为错题」等场景）。"""
    title: str = Field(default="", max_length=200)
    content: str = Field(default="", max_length=500_000)
    knowledge_tags: list[str] = Field(default_factory=list, max_length=100)
    subject: str = Field(default="", max_length=64)
    grade: str = Field(default="", max_length=64)
    question_ids: list[str] = Field(default_factory=list, max_length=200)


class NotePatchRequest(BaseModel):
    title: str | None = Field(default=None, max_length=200)
    content: str | None = Field(default=None, max_length=500_000)
    knowledge_tags: list[str] | None = Field(default=None, max_length=100)
    question_ids: list[str] | None = Field(default=None, max_length=200)
    sort_order: int | None = Field(default=None, ge=-1_000_000, le=1_000_000)

    model_config = ConfigDict(extra="forbid")


class NoteChatRequest(BaseModel):
    message: str = Field(min_length=1, max_length=50_000)


def _safe_download_filename(title: str) -> str:
    """把笔记标题清洗成安全的 Content-Disposition 文件名，防止换行注入响应头。"""
    name = re.sub(r"[\r\n\"\\/:*?<>|]+", "_", str(title or "").strip())
    name = name.strip(" ._")
    return (name or "note")[:120]


def _content_disposition_attachment(title: str, extension: str = ".md") -> str:
    """构造含非 ASCII 文件名的 Content-Disposition。

    Starlette 响应头按 latin-1 编码，直接放中文会抛 UnicodeEncodeError（500）。
    按 RFC 5987 提供 filename*（UTF-8 百分号编码），并附 ASCII 回退 filename。
    """
    from urllib.parse import quote
    name = _safe_download_filename(title)
    ascii_fallback = name.encode("ascii", "ignore").decode("ascii").strip(" ._") or "note"
    return f"attachment; filename=\"{ascii_fallback}{extension}\"; filename*=UTF-8''{quote(name)}{extension}"


def _merge_unique(existing, incoming, limit: int | None = None) -> list:
    merged = list(existing or [])
    for item in incoming or []:
        if item not in merged:
            merged.append(item)
            if limit is not None and len(merged) >= limit:
                break
    return merged[:limit] if limit is not None else merged


def _store_note_images(note_id: str, images: list[str], existing_count: int = 0) -> list[str]:
    """Decode validated uploads to storage instead of keeping large base64 blobs in SQLite."""
    if not images or existing_count >= MAX_IMAGES:
        return []
    folder = os.path.join(STORAGE_DIR, "notes", note_id)
    os.makedirs(folder, exist_ok=True)
    stored = []
    next_index = 0
    try:
        for encoded in images[:max(0, MAX_IMAGES - existing_count)]:
            payload = encoded.split(",", 1)[-1] if "," in encoded else encoded
            raw = base64.b64decode(payload, validate=True)
            if raw.startswith(b"\x89PNG\r\n\x1a\n"):
                ext = ".png"
            elif len(raw) >= 12 and raw[:4] == b"RIFF" and raw[8:12] == b"WEBP":
                ext = ".webp"
            elif raw.startswith(b"\xff\xd8\xff"):
                ext = ".jpg"
            else:
                raise ValueError("来源图片内容无效")
            while os.path.exists(os.path.join(folder, f"source_{next_index}{ext}")):
                next_index += 1
            filename = f"source_{next_index}{ext}"
            target = os.path.join(folder, filename)
            temp = target + ".tmp"
            try:
                with open(temp, "wb") as stream:
                    stream.write(raw)
                    stream.flush()
                    os.fsync(stream.fileno())
                os.replace(temp, target)
            finally:
                if os.path.exists(temp):
                    try:
                        os.remove(temp)
                    except OSError:
                        logger.warning("Failed to remove temporary note image: %s", temp)
            stored.append(f"/storage/notes/{note_id}/{filename}")
            next_index += 1
    except Exception:
        # 中途失败（磁盘满/编码异常）时清理本次已落盘文件，避免孤儿图片
        _remove_stored_note_images(note_id, stored)
        raise
    return stored


def _remove_stored_note_images(note_id: str, urls: list[str]) -> None:
    folder = os.path.abspath(os.path.join(STORAGE_DIR, "notes", note_id))
    for url in urls:
        filename = os.path.basename(str(url or ""))
        if not filename:
            continue
        target = os.path.abspath(os.path.join(folder, filename))
        if os.path.dirname(target) != folder:
            continue
        try:
            if os.path.isfile(target):
                os.remove(target)
        except OSError:
            logger.warning("Failed to roll back note image: %s", target)
    try:
        if os.path.isdir(folder) and not os.listdir(folder):
            os.rmdir(folder)
    except OSError:
        pass


def _remove_note_image_urls(urls: list[str]) -> None:
    root = os.path.abspath(os.path.join(STORAGE_DIR, "notes"))
    touched_folders = set()
    for url in urls:
        prefix = "/storage/notes/"
        value = str(url or "")
        if not value.startswith(prefix):
            continue
        relative = value[len(prefix):].replace("/", os.sep)
        target = os.path.abspath(os.path.join(root, relative))
        try:
            if os.path.commonpath([root, target]) != root:
                continue
        except ValueError:
            continue
        try:
            if os.path.isfile(target):
                os.remove(target)
                touched_folders.add(os.path.dirname(target))
        except OSError:
            logger.warning("Failed to remove note image: %s", target)
    for folder in touched_folders:
        try:
            if os.path.isdir(folder) and not os.listdir(folder):
                os.rmdir(folder)
        except OSError:
            pass


async def _refresh_auto_banks(db: AsyncSession):
    try:
        from config import ENABLE_NOTE_REFERENCES, get_setting
        from services.note_service import auto_create_banks_from_notes
        if ENABLE_NOTE_REFERENCES:
            threshold = max(1, int(get_setting("auto_bank_threshold", 5)))
            await auto_create_banks_from_notes(db, threshold)
    except Exception as exc:
        logger.warning("auto_create_banks_from_notes failed: %s", exc)


async def _extract_and_add_to_graph(note_id: str, title: str, content: str, tags: list) -> dict:
    """从笔记内容提取知识点并增量更新知识图谱。"""
    try:
        from services.knowledge_graph import extract_knowledge_from_text, add_note_to_graph
        extracted = await extract_knowledge_from_text(content, "")
        if extracted.get("concepts"):
            summary = add_note_to_graph(note_id, title, content, tags, extracted)
            return {"updated": True, "concepts": len(extracted.get("concepts", [])), **summary}
        return {"updated": False, "reason": "no_concepts"}
    except Exception as e:
        logger.warning("Knowledge graph update failed for note %s: %s", note_id, e)
        return {"updated": False, "error": str(e)}


async def _persist_structured_note(result: dict, *, source_type: str,
                                   fallback_title: str = "未命名笔记",
                                   source_images: list[str] | None = None) -> dict:
    """Persist an AI-structured note using one merge policy for every upload type."""
    from services.note_service import merge_notes

    def _as_text(value, default: str = "") -> str:
        """AI 可能返回数字/布尔/null 等非字符串字段，统一转成文本并截断，避免 .strip() 崩溃。"""
        if value is None:
            return default
        if isinstance(value, (int, float, bool)):
            return str(value)
        if isinstance(value, str):
            return value.strip() or default
        return default

    normalized = dict(result or {})
    normalized["knowledge_tags"] = canonicalize_tags(normalized.get("knowledge_tags", []))
    normalized["title"] = _as_text(normalized.get("title"), fallback_title)[:200]
    normalized["content"] = _as_text(normalized.get("content"))[:500_000]
    normalized["subject"] = _as_text(normalized.get("subject"))[:64]
    normalized["grade"] = _as_text(normalized.get("grade"))[:64]

    # 串行化整个 读快照→合并→写回 流程，防止并发上传基于过期快照互相覆盖
    async with _persist_note_lock:
        return await _persist_structured_note_locked(normalized, source_type=source_type,
                                                     source_images=source_images)


async def _persist_structured_note_locked(normalized: dict, *, source_type: str,
                                          source_images: list[str] | None) -> dict:
    """在持有 _persist_note_lock 的前提下执行入库/合并。"""
    from services.note_service import merge_notes

    async with async_session() as db:
        existing = await db.execute(select(Note).order_by(Note.updated_at.desc()))
        existing_notes = existing.scalars().all()
        existing_dicts = [
            {"id": n.id, "subject": n.subject or "", "title": n.title or "",
             "knowledge_tags": n.knowledge_tags or [], "content": n.content or ""}
            for n in existing_notes
        ]
        merge_idx, merged_content = merge_notes(normalized, existing_dicts)
        if 0 <= merge_idx < len(existing_notes):
            target = existing_notes[merge_idx]
            existing_sources = list(target.source_images or [])
            stored_sources = _store_note_images(target.id, source_images or [], len(existing_sources))
            target.content = merged_content or normalized["content"]
            target.knowledge_tags = canonicalize_tags(
                list(set(target.knowledge_tags or []) | set(normalized["knowledge_tags"]))
            )
            target.title = target.title or normalized["title"]
            target.subject = target.subject or normalized["subject"]
            target.grade = target.grade or normalized["grade"]
            target.source_images = _merge_unique(existing_sources, stored_sources, MAX_IMAGES)
            target.typical_questions = _merge_unique(
                target.typical_questions or [], normalized.get("typical_questions", []), 20
            )
            target.question_ids = _merge_unique(target.question_ids or [], normalized.get("question_ids", []), 200)
            target.references = _merge_unique(target.references or [], normalized.get("references", []), 200)
            if target.diagram_spec is None and normalized.get("diagram_spec") is not None:
                target.diagram_spec = normalized["diagram_spec"]
            target.is_structured = True
            target.updated_at = _utcnow()
            try:
                await db.commit()
            except Exception:
                await db.rollback()
                _remove_stored_note_images(target.id, stored_sources)
                raise
            await _refresh_auto_banks(db)
            # 知识图谱增量更新
            kg_result = await _extract_and_add_to_graph(target.id, target.title, merged_content or normalized["content"], target.knowledge_tags)
            return {"action": "merged", "target_id": target.id,
                    "note": {"id": target.id, "title": target.title,
                             "subject": target.subject, "knowledge_tags": target.knowledge_tags,
                             "is_structured": True, "knowledge_graph": kg_result}}

        note_id = gen_id()
        stored_sources = _store_note_images(note_id, source_images or [])
        note = Note(
            id=note_id, subject=normalized["subject"], grade=normalized["grade"],
            knowledge_tags=normalized["knowledge_tags"], title=normalized["title"],
            content=normalized["content"],
            typical_questions=normalized.get("typical_questions", []),
            source_images=stored_sources,
            question_ids=normalized.get("question_ids", []),
            references=normalized.get("references", []),
            diagram_spec=normalized.get("diagram_spec"),
            is_structured=True, source_type=source_type, auto_generated=False,
        )
        db.add(note)
        try:
            await db.commit()
        except Exception:
            await db.rollback()
            _remove_stored_note_images(note_id, stored_sources)
            raise
        await _refresh_auto_banks(db)
        # 知识图谱增量更新
        kg_result = await _extract_and_add_to_graph(note.id, note.title, normalized["content"], normalized["knowledge_tags"])
        return {"action": "created", "target_id": note.id,
                "note": {"id": note.id, "title": note.title,
                         "subject": note.subject, "knowledge_tags": note.knowledge_tags,
                         "is_structured": True, "knowledge_graph": kg_result}}
@router.post("")
async def create_note(req: NoteCreateRequest, db: AsyncSession = Depends(get_db)):
    """手动创建一条笔记。供批改报告「另存为错题」等入口调用。"""
    n = Note(
        id=gen_id(),
        title=(req.title or "").strip() or "未命名笔记",
        content=req.content or "",
        knowledge_tags=req.knowledge_tags or [],
        subject=req.subject or "",
        grade=req.grade or "",
        question_ids=req.question_ids or [],
        source_type="manual",
        is_structured=False,
        auto_generated=False,
    )
    db.add(n)
    await db.commit()
    return {"id": n.id, "message": "已创建笔记"}


@router.post("/upload-images")
async def api_upload_note_images(data: NoteUploadRequest):
    """上传图片，OCR识别 + AI分类 + 自动整合"""
    from services.note_service import ocr_image_async, classify_and_structure_async

    if not data.images:
        raise HTTPException(400, "至少需要一张图片")

    # 1. OCR 所有图片（异步）
    semaphore = asyncio.Semaphore(3)

    async def _ocr_note_image(img: str):
        async with semaphore:
            b64 = img.split(",")[-1] if "," in img else img
            return await ocr_image_async(b64)

    raw_results = await asyncio.gather(
        *(_ocr_note_image(img) for img in data.images), return_exceptions=True
    )
    all_texts = []
    failed_count = 0
    for item in raw_results:
        if isinstance(item, Exception):
            failed_count += 1
            logger.warning("Note image OCR failed: %s", item)
        elif item:
            all_texts.append(item)

    if not all_texts:
        raise HTTPException(400, "OCR 未能识别任何文字")

    combined_text = "\n\n".join(all_texts)

    # 2. AI 分类 + 结构化（异步）
    try:
        result = await classify_and_structure_async(combined_text, data.subject_hint)
    except HTTPException:
        raise
    except Exception as exc:
        log_error("notes.upload_images", f"AI structure failed: {exc}")
        _msg = str(exc).lower()
        if "401" in _msg or "authentication" in _msg or "api key" in _msg or "illegal header" in _msg:
            raise HTTPException(503, "AI 服务认证失败，请检查模型 API Key 配置")
        raise HTTPException(502, "笔记 AI 整理失败，请稍后重试")
    if not result.get("content"):
        result["content"] = combined_text

    persisted = await _persist_structured_note(
        result, source_type="image", fallback_title="图片笔记", source_images=data.images
    )
    persisted["processed_images"] = len(all_texts)
    persisted["failed_images"] = failed_count
    return persisted


@router.post("/upload-text")
async def api_upload_note_text(file: UploadFile = File(...), subject_hint: str = Form("", max_length=64)):
    """上传 txt/Markdown/Word/PDF 笔记，复用现有 AI 整理链路。"""
    filename = (file.filename or "").lower()
    if not filename.endswith((".txt", ".md", ".markdown", ".docx", ".pdf")):
        raise HTTPException(400, "仅支持 TXT、Markdown、DOCX 或 PDF 文件")
    raw = await file.read()
    max_size = 10 * 1024 * 1024 if filename.endswith((".docx", ".pdf")) else 2 * 1024 * 1024
    if len(raw) > max_size:
        raise HTTPException(413, f"笔记文件不能超过 {max_size // 1024 // 1024}MB")
    if not raw:
        raise HTTPException(400, "笔记文件为空")
    if filename.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + "\t".join(cell.text.strip() for cell in row.cells)
        except Exception as exc:
            raise HTTPException(400, f"Word 文档解析失败：{exc}")
    elif filename.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if len(reader.pages) > 100:
                raise HTTPException(400, "PDF 最多支持 100 页")
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except HTTPException:
            raise
        except ImportError:
            raise HTTPException(503, "PDF 解析组件未安装，请安装 requirements 后重试")
        except Exception as exc:
            raise HTTPException(400, f"PDF 解析失败：{exc}")
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030", errors="replace")
    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise HTTPException(400, "未从文档中提取到文字；扫描版 PDF 请改为上传图片")

    from services.note_service import classify_and_structure_async
    try:
        result = await classify_and_structure_async(text, subject_hint.strip())
    except HTTPException:
        raise
    except Exception as exc:
        log_error("notes.upload_text", f"AI structure failed: {exc}")
        _msg = str(exc).lower()
        if "401" in _msg or "authentication" in _msg or "api key" in _msg or "illegal header" in _msg:
            raise HTTPException(503, "AI 服务认证失败，请检查模型 API Key 配置")
        raise HTTPException(502, "笔记 AI 整理失败，请稍后重试")
    result["content"] = result.get("content") or text
    result["subject"] = result.get("subject") or subject_hint.strip()
    return await _persist_structured_note(
        result, source_type="text", fallback_title=filename.rsplit(".", 1)[0]
    )


@router.get("/knowledge-tree")
async def get_knowledge_tree(db: AsyncSession = Depends(get_db)):
    """按学科、知识点和笔记生成轻量知识树，供前端展示和夜间缓存。"""
    from services.note_service import build_knowledge_tree
    result = await db.execute(select(Note).order_by(Note.updated_at.desc()))
    notes = result.scalars().all()
    return build_knowledge_tree(notes)


@router.get("/{note_id}/download")
async def api_download_note(note_id: str):
    """下载笔记为 Markdown 文件"""
    from services.note_service import render_note_markdown
    from fastapi.responses import PlainTextResponse

    async with async_session() as db:
        note = await db.get(Note, note_id)
        if not note:
            raise HTTPException(404, "笔记不存在")

        md_text = render_note_markdown({
            "id": note.id,
            "title": note.title,
            "subject": note.subject,
            "grade": note.grade,
            "knowledge_tags": note.knowledge_tags or [],
            "content": note.content or "",
            "typical_questions": note.typical_questions or [],
            "diagram_spec": note.diagram_spec,
        })
        return PlainTextResponse(
            content=md_text,
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": _content_disposition_attachment(note.title)}
        )


@router.get("/search")
async def search_notes(keyword: str = "", type: str = "all", db: AsyncSession = Depends(get_db)):
    """搜索笔记、题目和题库"""
    kw = (keyword or "").strip()[:200]
    result = {"notes": [], "questions": [], "banks": []}
    if not kw:
        return result

    search_type = (type or "all").lower()
    # LIKE 通配符转义：用户输入的 % _ \ 必须按字面匹配
    escaped = kw.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")

    if search_type in ("note", "all"):
        r = await db.execute(
            select(Note).where(or_(Note.title.contains(escaped, escape="\\"),
                                   Note.content.contains(escaped, escape="\\"))).limit(50)
        )
        for n in r.scalars().all():
            note_refs = getattr(n, "references", None) or []
            resolved = resolve_note_references(n.content or "", note_refs)
            result["notes"].append({
                "id": n.id, "subject": n.subject, "grade": n.grade,
                "knowledge_tags": n.knowledge_tags, "title": n.title,
                "content": resolved[:200],
                "references": note_refs,
                "is_structured": n.is_structured,
                "updated_at": n.updated_at.isoformat() if n.updated_at else "",
            })

    if search_type in ("question", "all"):
        r = await db.execute(
            select(Question).where(
                or_(Question.ocr_text.contains(escaped, escape="\\"),
                    Question.question_html.contains(escaped, escape="\\"))
            ).limit(50)
        )
        for q in r.scalars().all():
            result["questions"].append({
                "id": q.id, "subject": q.subject, "grade": q.grade,
                "bank": q.bank,
                "knowledge_tags": q.knowledge_tags,
                "ocr_text": (q.ocr_text or "")[:200],
                "updated_at": q.updated_at.isoformat() if q.updated_at else "",
            })

    if search_type in ("bank", "all"):
        rows = await db.execute(
            select(Question.bank, func.count(Question.id))
            .where(Question.bank.contains(escaped, escape="\\"))
            .group_by(Question.bank)
        )
        for name, count in rows.all():
            if name:
                result["banks"].append({"name": name, "count": count})

    return result


@router.post("/auto-organize")
async def api_auto_organize():
    """自动整理所有笔记：去重合并（HTTP 端点，无额外参数）"""
    return await _do_auto_organize()


async def _do_auto_organize(classic_models_context: list[str] = None):
    """内部实现：支持注入经典模型上下文（供夜间巡逻直接调用）"""
    from services.note_service import merge_notes

    # 与单篇入库/合并共用同一把锁，避免自动整理删除的笔记
    # 与并发上传写入的目标互相覆盖
    async with _persist_note_lock:
        return await _do_auto_organize_locked(classic_models_context)


async def _do_auto_organize_locked(classic_models_context: list[str] = None) -> dict:
    from services.note_service import merge_notes

    async with async_session() as db:
        result = await db.execute(select(Note).order_by(Note.updated_at.desc()))
        notes = result.scalars().all()

        merged_count = 0
        kept_ids = set()
        removed_ids = set()

        for i, note_a in enumerate(notes):
            if note_a.id in kept_ids or note_a.id in removed_ids:
                continue
            kept_ids.add(note_a.id)
            for j, note_b in enumerate(notes):
                if i >= j or note_b.id in kept_ids or note_b.id in removed_ids:
                    continue
                merge_idx, merged = merge_notes(
                    {"subject": note_a.subject or "", "knowledge_tags": note_a.knowledge_tags or [], "content": note_a.content or ""},
                    [{"id": note_b.id, "subject": note_b.subject or "", "knowledge_tags": note_b.knowledge_tags or [], "content": note_b.content or ""}],
                    overlap_threshold=0.6,
                )
                if merge_idx >= 0 and merged:
                    note_a.content = merged
                    existing_tags = set(note_a.knowledge_tags or [])
                    new_tags = set(note_b.knowledge_tags or [])
                    note_a.knowledge_tags = canonicalize_tags(list(existing_tags | new_tags))
                    # 合并后保留缺失的标题/学科/学段与来源类型语义
                    if not note_a.title and note_b.title:
                        note_a.title = note_b.title
                    if not note_a.subject and note_b.subject:
                        note_a.subject = note_b.subject
                    if not note_a.grade and note_b.grade:
                        note_a.grade = note_b.grade
                    src_a = getattr(note_a, "source_type", "") or ""
                    src_b = getattr(note_b, "source_type", "") or ""
                    if src_b and src_a != src_b:
                        note_a.source_type = f"{src_a}+{src_b}" if src_a else src_b
                    note_a.source_images = _merge_unique(note_a.source_images or [], note_b.source_images or [], MAX_IMAGES)
                    note_a.typical_questions = _merge_unique(note_a.typical_questions or [], note_b.typical_questions or [], 20)
                    note_a.question_ids = _merge_unique(note_a.question_ids or [], note_b.question_ids or [], 200)
                    note_a.references = _merge_unique(note_a.references or [], note_b.references or [], 200)
                    if note_a.diagram_spec is None and note_b.diagram_spec is not None:
                        note_a.diagram_spec = note_b.diagram_spec
                    note_a.updated_at = _utcnow()
                    await db.delete(note_b)
                    removed_ids.add(note_b.id)
                    merged_count += 1

        await db.commit()

    # 合并后清理被删笔记在知识图谱中的贡献（节点被多篇笔记共享时仅降权重）
    if removed_ids:
        try:
            from services.knowledge_graph import remove_note_from_graph
            for rid in removed_ids:
                remove_note_from_graph(rid)
        except Exception as exc:
            logger.warning("Knowledge graph cleanup after auto-organize failed: %s", exc)
    return {"message": f"整理完成，合并了 {merged_count} 篇笔记"}


@router.put("/{note_id}/full")
async def update_note_full(note_id: str, data: NoteUpdateRequest):
    """完整修改笔记（含diagram_spec）"""
    async with async_session() as db:
        note = await db.get(Note, note_id)
        if not note:
            raise HTTPException(404, "笔记不存在")
        fields = data.model_fields_set
        if "title" in fields:
            note.title = data.title or ""
        if "content" in fields:
            note.content = data.content or ""
        if "knowledge_tags" in fields:
            note.knowledge_tags = canonicalize_tags(data.knowledge_tags or [])
        if "subject" in fields:
            note.subject = data.subject or ""
        if "grade" in fields:
            note.grade = data.grade or ""
        if "diagram_spec" in fields:
            note.diagram_spec = data.diagram_spec
        note.updated_at = _utcnow()
        await db.commit()
        # 正文变化后同步刷新图谱摘要片段，避免详情页展示旧内容
        if "content" in fields:
            try:
                from services.knowledge_graph import refresh_note_snippet
                refresh_note_snippet(note_id, note.content or "")
            except Exception as exc:
                logger.warning("Knowledge graph snippet refresh failed for %s: %s", note_id, exc)
        return {"message": "已更新", "id": note_id}


@router.post("/{note_id}/references")
async def update_note_references(note_id: str, req: NoteReferencesRequest, db: AsyncSession = Depends(get_db)):
    """更新笔记引用，并将缺失的引用标记同步插入 content"""
    note = await db.get(Note, note_id)
    if not note:
        raise HTTPException(404, "笔记不存在")

    refs = []
    for r in req.references:
        if not isinstance(r, dict):
            continue
        ref_type = str(r.get("type", "")).strip().lower()
        if ref_type not in ("question", "bank"):
            continue
        ref_id = str(r.get("id", "")).strip()[:64]
        # 清洗 name：去掉会破坏 [[BANK:...]] 标记的括号与换行，并限制长度
        ref_name = str(r.get("name", "")).strip()[:100]
        ref_name = ref_name.replace("[[", "").replace("]]", "").replace("\n", " ").replace("\r", " ")
        if ref_type == "question" and not ref_id:
            continue
        if ref_type == "bank" and not ref_name:
            continue
        refs.append({"type": ref_type, "id": ref_id, "name": ref_name})

    question_ref_ids = list(dict.fromkeys(ref["id"] for ref in refs if ref["type"] == "question"))
    if question_ref_ids:
        result = await db.execute(select(Question).where(Question.id.in_(question_ref_ids)))
        q_map = {q.id: q for q in result.scalars().all()}
        from services.question_challenge import has_unresolved_high_challenge
        invalid = [
            qid for qid in question_ref_ids
            if qid not in q_map or q_map[qid].status != "done"
            or q_map[qid].source_type in ("search_query", "correction_query")
            or has_unresolved_high_challenge(q_map[qid].audit_flags, q_map[qid].is_resolved)
        ]
        if invalid:
            raise HTTPException(400, detail="引用中存在不可用题目：" + ", ".join(invalid[:10]))

    note.references = refs

    content = getattr(note, "content", None) or ""
    markers_to_add = []
    for ref in refs:
        if ref["type"] == "question":
            marker = f"[[QUESTION:{ref['id']}]]"
        else:
            marker = f"[[BANK:{ref['name']}]]"
        if marker not in content:
            markers_to_add.append(marker)

    if markers_to_add:
        sep = "\n\n" if content.strip() else ""
        note.content = content.rstrip() + sep + "\n".join(markers_to_add)

    note.updated_at = _utcnow()
    await db.commit()
    return {"message": "已更新", "references": getattr(note, "references", None) or []}


@router.get("")
async def list_notes(subject: str = "", grade: str = "",
                     tag: str = "", db: AsyncSession = Depends(get_db)):
    """返回笔记列表，按学科分组"""
    q = select(Note).order_by(Note.sort_order.desc(), Note.created_at.desc())
    if subject: q = q.where(Note.subject == subject)
    if grade: q = q.where(Note.grade == grade)
    if tag:
        escaped_tag = str(tag).replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_").replace('"', '')
        q = q.where(Note.knowledge_tags.cast(TEXT).like(f'%"{escaped_tag}"%', escape="\\"))
    r = await db.execute(q)
    notes = r.scalars().all()

    def _note_item(n: Note) -> dict:
        note_refs = getattr(n, "references", None) or []
        resolved = resolve_note_references(n.content or "", note_refs)
        return {
            "id": n.id, "subject": n.subject, "grade": n.grade,
            "knowledge_tags": n.knowledge_tags, "title": n.title,
            "content": resolved[:200],
            "references": note_refs,
            "question_count": len(n.question_ids or []),
            "sort_order": n.sort_order, "auto_generated": n.auto_generated,
            "has_diagram": n.diagram_spec is not None,
            "is_structured": n.is_structured,
            "created_at": n.created_at.isoformat() if n.created_at else "",
            "updated_at": n.updated_at.isoformat() if n.updated_at else "",
        }

    # 按subject分组
    groups = {}
    for n in notes:
        s = n.subject or "未分类"
        groups.setdefault(s, []).append(_note_item(n))

    return {"notes": [_note_item(n) for n in notes], "groups": groups}


@router.get("/{note_id}")
async def get_note(note_id: str, db: AsyncSession = Depends(get_db)):
    n = await db.get(Note, note_id)
    if not n: raise HTTPException(404, detail="笔记不存在")
    qids = n.question_ids or []
    questions = []
    if qids:
        async with async_session() as _db:
            result = await _db.execute(select(Question).where(Question.id.in_(qids[:30])))
            q_map = {q.id: q for q in result.scalars().all()}
            for qid in qids[:30]:
                q = q_map.get(qid)
                if q:
                    questions.append({
                        "id": q.id, "subject": q.subject, "grade": q.grade,
                        "ocr_text": (q.ocr_text or "")[:80],
                        "knowledge_tags": q.knowledge_tags
                    })
    note_refs = getattr(n, "references", None) or []
    return {
        "id": n.id, "subject": n.subject, "grade": n.grade,
        "knowledge_tags": n.knowledge_tags, "title": n.title,
        "content": resolve_note_references(n.content or "", note_refs),
        "question_ids": n.question_ids,
        "questions": questions, "auto_generated": n.auto_generated,
        "typical_questions": n.typical_questions or [],
        "source_images": n.source_images or [],
        "diagram_spec": n.diagram_spec,
        "is_structured": n.is_structured,
        "source_type": n.source_type,
        "references": note_refs,
        "created_at": n.created_at.isoformat() if n.created_at else "",
        "updated_at": n.updated_at.isoformat() if n.updated_at else "",
    }


@router.put("/{note_id}")
async def update_note(note_id: str, data: NotePatchRequest, db: AsyncSession = Depends(get_db)):
    n = await db.get(Note, note_id)
    if not n: raise HTTPException(404, detail="笔记不存在")
    values = data.model_dump(exclude_unset=True)
    if "knowledge_tags" in values:
        values["knowledge_tags"] = canonicalize_tags(values["knowledge_tags"] or [])
    if "question_ids" in values:
        ids = list(dict.fromkeys(values["question_ids"] or []))
        if ids:
            result = await db.execute(select(Question.id).where(
                Question.id.in_(ids),
                Question.status == "done",
                or_(Question.source_type.is_(None),
                    ~Question.source_type.in_(["search_query", "correction_query"])),
                or_(Question.is_resolved.is_(True), Question.audit_flags.is_(None),
                    ~Question.audit_flags.cast(TEXT).ilike("%question_challenge_high%")),
            ))
            existing = {row[0] for row in result.fetchall()}
            missing = [qid for qid in ids if qid not in existing]
            if missing:
                raise HTTPException(400, detail="关联题目不存在或未完成：" + ", ".join(missing[:10]))
        values["question_ids"] = ids
    for key, value in values.items():
        setattr(n, key, value)
    n.updated_at = _utcnow()
    await db.commit()
    if "content" in values:
        try:
            from services.knowledge_graph import refresh_note_snippet
            refresh_note_snippet(note_id, n.content or "")
        except Exception as exc:
            logger.warning("Knowledge graph snippet refresh failed for %s: %s", note_id, exc)
    return {"message": "已更新"}


@router.delete("/{note_id}")
async def delete_note(note_id: str, db: AsyncSession = Depends(get_db)):
    n = await db.get(Note, note_id)
    if not n: raise HTTPException(404, detail="笔记不存在")
    candidate_urls = list(n.source_images or [])
    await db.delete(n)
    await db.commit()
    # 同步清理知识图谱中该笔记的贡献，避免残留引用和虚高权重
    try:
        from services.knowledge_graph import remove_note_from_graph
        remove_note_from_graph(note_id)
    except Exception as exc:
        logger.warning("Knowledge graph cleanup failed for note %s: %s", note_id, exc)
    if candidate_urls:
        result = await db.execute(select(Note.source_images))
        referenced = {
            str(url)
            for (items,) in result.fetchall()
            for url in (items or [])
        }
        _remove_note_image_urls([url for url in candidate_urls if str(url) not in referenced])
    return {"message": "已删除"}


@router.post("/ai-generate")
async def ai_generate_notes():
    """LLM 自动整理笔记：让DeepSeek分析所有题目，归纳考点并撰写总结"""
    from services.ai_service import ai_service

    # 1. 获取所有已完成题目
    async with async_session() as _db:
        r = await _db.execute(select(Question).where(
            Question.status == "done",
            or_(Question.source_type.is_(None),
                ~Question.source_type.in_(["search_query", "correction_query"])),
            or_(Question.is_resolved.is_(True), Question.audit_flags.is_(None),
                ~Question.audit_flags.cast(TEXT).ilike("%question_challenge_high%")),
        ))
        all_qs = r.scalars().all()

    if not all_qs:
        return {"message": "暂无已完成题目", "note_count": 0}

    # 2. 构建题目摘要发给LLM
    q_summaries = []
    for q in all_qs:
        tags_str = ", ".join(q.knowledge_tags or [])
        ocr = (q.ocr_text or q.question_html or "")[:100].replace("\n", " ")
        sa = (q.standard_answer or "")[:50]
        q_summaries.append(f"[{q.subject}][{q.grade}] {ocr}... 标答:{sa} 标签:{tags_str}")

    # 按学科分组避免一次发送太多
    by_subject = {}
    for q in all_qs:
        s = q.subject or "未分类"
        by_subject.setdefault(s, []).append(q)

    all_notes = []
    expected_batches = sum((len(qs) + 29) // 30 for qs in by_subject.values())
    successful_batches = 0
    batch_failures = []

    for subject, qs in by_subject.items():
        # 每批最多30题
        for batch_start in range(0, len(qs), 30):
            batch = qs[batch_start:batch_start+30]
            batch_text = []
            for q in batch:
                ocr = (q.ocr_text or q.question_html or "")[:120].replace("\n", " ")
                sa = (q.standard_answer or "")[:60]
                at = (q.answer_html or "")[:120].replace("\n", " ").replace("<!-- SCORE_SPLIT -->", " | ")
                batch_text.append(f"ID:{q.id[:8]} [{q.grade}][题{ocr}][答案{sa}][解法{at}]")

            prompt = (
                f"请分析以下{subject}题目，**忽略已有标签**，凭真正的解题思路和几何/代数结构，归纳出3-8个'解题套路/模型'，并以JSON格式输出。\n\n"
                f"题目列表：\n" + "\n".join(batch_text) + "\n\n"
                f"【什么是解题套路/模型】\n"
                f"- 不是知识点标签（如'全等三角形'），而是具体的解题模型和手法：\n"
                f"  例如：'一线三等角模型'（一条直线上出现三个等角则蕴含全等/相似）\n"
                f"  例如：'手拉手旋转模型'（两个等腰三角形共顶点旋转得全等）\n"
                f"  例如：'角平分线+平行出等腰'（平行线与角平分线交于底边即得等腰三角形）\n"
                f"  例如：'含30°直角三角形边比1:√3:2'\n"
                f"  例如：'二次函数求最值配方法'\n"
                f"  例如：'串联分压/并联分流'（物理电路分析套路）\n"
                f"- 每个套路必须包含：①核心识别标志（看到什么条件就知道用这个套路）②解题步骤（123步）③易错点\n\n"
                f"【内容格式要求】\n"
                f"使用Markdown格式书写content：\n"
                f"- 用 ## 表示套路名称\n"
                f"- 用 **粗体** 强调关键结论\n"
                f"- 有序列表1. 2. 3. 表示解题步骤\n"
                f"- 用 `$...$` 包裹数学公式\n"
                f"- 用 > 引用标记易错点\n\n"
                f"输出JSON（不要markdown包裹）：\n"
                f'{{"notes":[\n'
                f'  {{"title":"套路名称（如：一线三等角模型）","content":"## 识别标志\\\\n...\\\\n## 解题步骤\\\\n...\\\\n> **易错点：**...","tags":["相关知识点"],"question_ids":["题目ID前缀..."]}},\n'
                f'  ...\n'
                f']}}'
            )
            try:
                result = await ai_service.deepseek_json(
                    [{"role": "user", "content": prompt}], max_tokens=8192, scope="notes_classify"
                )
                generated_items = result.get("notes", []) if isinstance(result, dict) else []
                if not isinstance(generated_items, list) or not generated_items:
                    raise ValueError("模型未返回有效笔记列表")
                successful_batches += 1
                for item in generated_items:
                    if not isinstance(item, dict):
                        continue
                    # 解析题目ID（前缀匹配）；qids 非列表或含非字符串元素时逐项防御
                    qids = item.get("question_ids", [])
                    if not isinstance(qids, (list, tuple)):
                        qids = []
                    resolved_ids = []
                    for prefix in qids:
                        prefix = str(prefix or "").strip()
                        if not prefix:
                            continue
                        for q in batch:
                            if q.id.startswith(prefix):
                                resolved_ids.append(q.id)
                                break
                    if resolved_ids:
                        title = str(item.get("title", "考点") or "考点").strip()[:200] or "考点"
                        content = str(item.get("content", "") or "").strip()[:200_000]
                        tags = canonicalize_tags(item.get("tags", []) if isinstance(item.get("tags"), list) else [])[:100]
                        all_notes.append({
                            "subject": subject,
                            "grade": str(item.get("grade", "") or "")[:64],
                            "title": title,
                            "content": content,
                            "tags": tags,
                            "question_ids": resolved_ids,
                            "sort_order": len(resolved_ids),  # 频次排序
                        })
            except Exception as e:
                logger.warning("LLM note generation failed for %s: %s", subject, e)
                batch_failures.append({"subject": subject, "batch_start": batch_start, "error": str(e)[:200]})

    if batch_failures or successful_batches != expected_batches or not all_notes:
        log_error(
            "notes.ai_generate",
            f"Generation incomplete: success={successful_batches}/{expected_batches}, failures={batch_failures[:5]}",
        )
        raise HTTPException(
            status_code=502,
            detail={
                "message": "AI 笔记生成未完整成功，已保留原有笔记",
                "successful_batches": successful_batches,
                "expected_batches": expected_batches,
                "failures": batch_failures,
            },
        )

    # 3. 写入数据库（先增量合并、再清理孤儿笔记，避免先删后创导致数据丢失）
    created = 0
    merged = 0
    async with async_session() as _db:
        # 查询现有 auto_generated 笔记
        r = await _db.execute(select(Note).where(Note.auto_generated == True))
        existing_auto = {(n.subject or "", n.title): n for n in r.scalars().all()}
        new_titles = set()
        for item in all_notes:
            title = item["title"]
            note_key = (item["subject"], title)
            new_titles.add(note_key)
            if note_key in existing_auto:
                # 更新已有笔记
                n = existing_auto[note_key]
                n.content = item["content"]
                n.knowledge_tags = item["tags"]
                n.question_ids = item["question_ids"]
                n.sort_order = item["sort_order"]
                n.subject = item["subject"]
                n.grade = item["grade"]
                n.updated_at = _utcnow()
                merged += 1
            else:
                # 新建笔记
                n = Note(
                    subject=item["subject"], grade=item["grade"],
                    knowledge_tags=item["tags"], title=title,
                    content=item["content"], question_ids=item["question_ids"],
                    sort_order=item["sort_order"], auto_generated=True
                )
                _db.add(n)
                created += 1
        # 清理不在新生成列表中的旧 auto_generated 笔记
        deleted = 0
        for old_key, old_note in existing_auto.items():
            if old_key not in new_titles:
                await _db.delete(old_note)
                deleted += 1
        await _db.commit()

    return {"message": f"AI 已整理完成：新建 {created} 篇，合并 {merged} 篇，清理 {deleted} 篇，覆盖 {len(all_qs)} 道题目"}


@router.post("/{note_id}/chat")
async def chat_note(note_id: str, req: NoteChatRequest):
    """与笔记对话"""
    from services.ai_service import ai_service
    async with async_session() as _db:
        n = await _db.get(Note, note_id)
        if not n: raise HTTPException(404, detail="笔记不存在")
        qids = n.question_ids or []
        note_content = n.content or ""
        if len(note_content) > 60_000:
            raise HTTPException(413, detail="笔记正文过长，暂不适合直接对话；请先拆分笔记")
        qs = []
        for qid in qids[:5]:
            q = await _db.get(Question, qid)
            if q: qs.append(f"- {(q.ocr_text or q.question_html or '')[:100]}... ({q.id[:8]})")
    qs_text = "\n".join(qs)
    prompt = (
        f"[考点] {n.title}\n"
        f"[学科/年级] {n.subject}/{n.grade}\n"
        f"[知识点] {', '.join(n.knowledge_tags or [])}\n"
        f"[笔记] {note_content}\n"
        f"[关联题目]\n{qs_text}\n\n"
        f"用户问：{req.message}\n"
        f"请基于笔记和相关题目回答。"
    )
    try:
        reply = await ai_service.deepseek_chat(
            [{"role": "user", "content": prompt}], max_tokens=4096, scope="chat"
        )
    except Exception as exc:
        log_error("notes.chat", f"AI chat failed for {note_id}: {exc}")
        _msg = str(exc).lower()
        if "401" in _msg or "authentication" in _msg or "api key" in _msg or "illegal header" in _msg:
            raise HTTPException(status_code=503, detail="AI 服务认证失败，请检查模型 API Key 配置")
        raise HTTPException(status_code=502, detail="AI 对话失败，请稍后重试")
    return {"reply": reply}
