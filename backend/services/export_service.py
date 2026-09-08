import os
import re
import asyncio
import tempfile
from concurrent.futures import ThreadPoolExecutor
from sqlalchemy.ext.asyncio import AsyncSession
from models.models import Paper
from models.database import async_session
from config import PAPERS_DIR
from logger import get_logger, log_error

logger = get_logger()

executor = ThreadPoolExecutor(max_workers=2)


class ExportService:

    async def export_html_word(self, paper_id: str, html: str, variant: str = "paper") -> str:
        """Export the exact requested HTML variant instead of always returning paper.docx."""
        if not re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", paper_id or ""):
            raise ValueError("非法的试卷 ID")
        if not re.fullmatch(r"[a-z0-9_-]{1,32}", variant or ""):
            raise ValueError("非法的 Word 导出类型")
        paper_dir = os.path.join(PAPERS_DIR, paper_id)
        os.makedirs(paper_dir, exist_ok=True)
        output_path = os.path.join(paper_dir, f"{variant}.docx")
        loop = asyncio.get_running_loop()
        await loop.run_in_executor(executor, self._html_to_word, html, output_path)
        if not os.path.isfile(output_path):
            raise RuntimeError("Word 文件未生成")
        return output_path

    async def export_word(self, paper_id: str, include_answer: bool = True) -> dict:
        if not re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", paper_id or ""):
            raise ValueError("非法的试卷 ID")
        async with async_session() as db:
            paper = await db.get(Paper, paper_id)
            if not paper:
                raise ValueError(f"Paper {paper_id} not found")

            paper_dir = os.path.join(PAPERS_DIR, paper_id)
            os.makedirs(paper_dir, exist_ok=True)
            result = {}

            loop = asyncio.get_running_loop()
            word_path = os.path.join(paper_dir, "paper.docx")
            generated_paths = []
            try:
                await loop.run_in_executor(executor, self._html_to_word, paper.paper_html, word_path)
                if not os.path.exists(word_path):
                    raise RuntimeError("Word 试卷文件未生成")
                generated_paths.append(word_path)
                paper.paper_word_path = word_path
                result["paper_word"] = word_path

                if include_answer and paper.answer_html:
                    answer_word = os.path.join(paper_dir, "answer.docx")
                    await loop.run_in_executor(executor, self._html_to_word, paper.answer_html, answer_word)
                    if not os.path.exists(answer_word):
                        raise RuntimeError("Word 答案文件未生成")
                    generated_paths.append(answer_word)
                    paper.answer_word_path = answer_word
                    result["answer_word"] = answer_word

                await db.commit()
                return result
            except Exception:
                # 半成品失败时清理已生成文件，避免残留伪路径误导下载
                for path in generated_paths:
                    try:
                        os.remove(path)
                    except OSError:
                        pass
                paper.paper_word_path = None
                paper.answer_word_path = None
                raise

    def _html_to_word(self, html: str, output_path: str):
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from bs4 import BeautifulSoup
        except ImportError as e:
            raise RuntimeError(f"Word 导出依赖缺失: {e}") from e

        try:
            doc = Document()
            style = doc.styles['Normal']
            style.font.size = Pt(12)

            body_match = re.search(r'<body[^>]*>(.*?)</body>', html, re.DOTALL | re.IGNORECASE)
            content = body_match.group(1) if body_match else html

            # Remove script and style tags before processing
            content = re.sub(r'<script[^>]*>.*?</script>', '', content, flags=re.DOTALL | re.IGNORECASE)
            content = re.sub(r'<style[^>]*>.*?</style>', '', content, flags=re.DOTALL | re.IGNORECASE)

            # Strip LaTeX/math delimiters for plain text in Word
            content = re.sub(r'\$\$([^$]+)\$\$', r'\1', content)
            content = re.sub(r'\$([^$]+)\$', r'\1', content)
            content = re.sub(r'\\\[([^\]]+)\\\]', r'\1', content)
            content = re.sub(r'\\\(([^\)]+)\\\)', r'\1', content)

            soup = BeautifulSoup(content, 'html.parser')

            block_names = ['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'div',
                           'section', 'article', 'blockquote', 'li', 'td', 'th']
            container_names = {'div', 'section', 'article', 'blockquote', 'td', 'th'}
            for tag in soup.find_all(block_names):
                # 只导出最内层内容块，避免父 div/section 与子 p 被重复写入 Word。
                if tag.name in container_names and tag.find(block_names, recursive=False):
                    continue
                text = tag.get_text(' ', strip=True)
                if not text:
                    continue
                if tag.name in ('h1', 'h2'):
                    doc.add_heading(text, level=1 if tag.name == 'h1' else 2)
                elif tag.name in ('h3', 'h4', 'h5', 'h6'):
                    doc.add_heading(text, level=3)
                elif tag.name == 'li':
                    doc.add_paragraph(text, style='List Bullet')
                else:
                    doc.add_paragraph(text)

            # 先写临时文件再原子替换，避免并发导出或进程中断留下损坏的 docx
            temp_path = ""
            try:
                fd = tempfile.NamedTemporaryFile(
                    mode="wb", dir=os.path.dirname(output_path) or ".",
                    prefix=".docx-", suffix=".tmp", delete=False
                )
                temp_path = fd.name
                fd.close()
                doc.save(temp_path)
                os.replace(temp_path, output_path)
            finally:
                if temp_path and os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except OSError:
                        pass
        except Exception as e:
            log_error("export", f"Word export failed: {e}")
            raise RuntimeError(f"Word 导出失败: {e}") from e


export_service = ExportService()
